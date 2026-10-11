"""Importiert die Ortsliste.

Zwei Quellen werden unterstützt:

* die aufbereitete CSV aus ``data/places/gemeinden.csv``
  (Spalten: ags,name,district_ags,district_name,state_code,lat,lon)
* der amtliche Auszug aus dem Gemeindeverzeichnis GV-ISys von DESTATIS
  (``AuszugGV...xlsx``, Blatt "Onlineprodukt_Gemeinden"), Option ``--gvisys``

Der Import ersetzt keine Orte, an denen Nutzer hängen. Verschwundene AGS
werden gemeldet; die Zuordnung alter auf neue Gemeinden gehört in
``place_merges`` (Befehl ``orte_fusion``).
"""

import csv
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from places.models import Place, normalize

SUFFIXES = (
    ", Landeshauptstadt", ", Stadt", ", Hansestadt", ", Universitätsstadt", ", Kreisstadt", ", Markt",
    ", Flecken", ", gemfr. Gebiet", ", Freie und Hansestadt", ", Wissenschaftsstadt", ", Große Kreisstadt",
    ", St.", ", M.", ", GKSt", ", Kurort", ", Bergstadt", ", Marktflecken", ", Landkreis", ", kreisfreie Stadt",
    ", Kreisfreie Stadt", ", Stadtkreis", ", documenta-Stadt", ", Klingenstadt", ", Lutherstadt",
    ", Barbarossastadt", ", Hochschulstadt", ", Inselgemeinde", ", Regionalverband",
)


def clean_name(name: str) -> str:
    name = " ".join(str(name).split())
    changed = True
    while changed:
        changed = False
        for suffix in SUFFIXES:
            if name.endswith(suffix):
                name = name[: -len(suffix)]
                changed = True
    return name


def read_csv(path: Path) -> list[dict]:
    with path.open(encoding="utf-8", newline="") as fh:
        rows = list(csv.DictReader(fh))
    need = {"ags", "name", "district_ags", "district_name", "state_code", "lat", "lon"}
    if not rows or not need.issubset(rows[0].keys()):
        raise CommandError(f"{path}: erwartete Spalten {sorted(need)}")
    return rows


def read_gvisys(path: Path) -> list[dict]:
    """Liest den DESTATIS-Auszug. Satzart 40 = Kreis, 60 = Gemeinde."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise CommandError("Für --gvisys wird openpyxl benötigt.") from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = next((wb[n] for n in wb.sheetnames if "Gemeinden" in n), wb[wb.sheetnames[-1]])
    districts: dict[str, str] = {}
    rows: list[dict] = []
    for raw in ws.iter_rows(values_only=True):
        cells = ["" if c is None else str(c).strip() for c in raw]
        if len(cells) < 16 or cells[0] not in {"40", "60"}:
            continue
        land, rb, kreis, gem, name = cells[2].zfill(2), cells[3] or "0", cells[4].zfill(2), cells[6].zfill(3), cells[7]
        district_ags = f"{land}{rb}{kreis}"
        if cells[0] == "40":
            districts[district_ags] = clean_name(name)
            continue
        try:
            lon = float(cells[14].replace(",", "."))
            lat = float(cells[15].replace(",", "."))
        except ValueError:
            continue
        rows.append({
            "ags": f"{district_ags}{gem}", "name": clean_name(name), "district_ags": district_ags,
            "district_name": "", "state_code": land, "lat": lat, "lon": lon,
        })
    for row in rows:
        row["district_name"] = districts.get(row["district_ags"], row["name"])
    return rows


class Command(BaseCommand):
    help = "Importiert die Ortsliste (CSV oder GV-ISys-Auszug von DESTATIS)."

    def add_arguments(self, parser):
        parser.add_argument("datei", type=Path)
        parser.add_argument("--gvisys", action="store_true", help="Datei ist der xlsx-Auszug aus GV-ISys")

    @transaction.atomic
    def handle(self, *args, datei: Path, gvisys: bool, **options):
        if not datei.exists():
            raise CommandError(f"{datei} nicht gefunden.")
        rows = read_gvisys(datei) if gvisys else read_csv(datei)
        if not rows:
            raise CommandError("Keine Orte gelesen.")
        seen: set[str] = set()
        places = []
        for row in rows:
            ags = str(row["ags"]).strip().zfill(8)
            lat, lon = float(row["lat"]), float(row["lon"])
            if ags in seen or len(ags) != 8 or not ags.isdigit():
                continue
            if not (47.0 <= lat <= 55.2 and 5.5 <= lon <= 15.3):
                self.stderr.write(f"Übersprungen (Koordinate außerhalb Deutschlands): {ags} {row['name']}")
                continue
            seen.add(ags)
            name = clean_name(row["name"])
            places.append(Place(
                ags=ags, name=name, search_name=normalize(name), district_ags=ags[:5],
                district_name=clean_name(row["district_name"]) or name, state_code=ags[:2], lat=lat, lon=lon,
            ))
        Place.objects.bulk_create(
            places, batch_size=2000, update_conflicts=True, unique_fields=["ags"],
            update_fields=["name", "search_name", "district_ags", "district_name", "state_code", "lat", "lon"],
        )
        gone = Place.objects.exclude(ags__in=seen)
        used = gone.filter(models_in_use())
        unused_count = gone.exclude(pk__in=used.values("pk")).delete()[0]
        self.stdout.write(self.style.SUCCESS(f"{len(places)} Orte importiert, {unused_count} nicht mehr vorhandene entfernt."))
        for place in used:
            self.stderr.write(
                f"AGS {place.ags} ({place.name}) fehlt im neuen Stand, wird aber von Nutzern verwendet. "
                "Bitte mit 'orte_fusion' auf die neue Gemeinde abbilden."
            )


def models_in_use():
    from django.db.models import Exists, OuterRef, Q

    from accounts.models import PlaceChangeRequest, User

    return (
        Exists(User.objects.filter(home_id=OuterRef("pk")))
        | Exists(User.objects.filter(residence_id=OuterRef("pk")))
        | Exists(PlaceChangeRequest.objects.filter(Q(new_place_id=OuterRef("pk")) | Q(old_place_id=OuterRef("pk"))))
    )
